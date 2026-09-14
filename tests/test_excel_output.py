"""Task 9.x: khuon cot day du + cau truc "moi loai san pham 1 sheet"."""
import json

from openpyxl import load_workbook

from crawler.record import (
    COLUMNS,
    UNGROUPED_SHEET,
    ProductRecord,
    import_legacy_xlsx,
    write_records_to_excel,
)
from crawler.record.price import LIEN_HE

# Hang header nguyen van cua sheet "2. LED Downlight" trong
# `product_Metadata (1).xlsx` - chep tay vao day de test khong phu thuoc vao
# viec file .xlsx goc co nam trong repo hay khong.
REFERENCE_HEADERS = [
    "STT", "Product_ID", "Tên sản phẩm", "Mã Sản Phẩm", "Mã SAP",
    " category 1 ", " category 2", " category 3", "Tags", "Giá",
    "Giá đối chiếu", "Link sản phẩm", "Link ảnh sản phẩm", "Link mua hàng online",
    "Tóm tắt TSKT", "Tóm tắt ưu điểm, tính năng", "Thông số kỹ thuật ",
    "Link file HDSD", "VD HDSD", "Nội dung Ưu điểm SP",
]


def _record(**kwargs) -> ProductRecord:
    base = dict(
        product_id="1", ten_san_pham="A", ma_san_pham="X1", category_1="Đèn LED âm trần",
        tags={"cong_suat": "10W"}, gia=166000.0,
        link_san_pham="https://x/1", link_anh_san_pham="https://x/1.png",
    )
    base.update(kwargs)
    return ProductRecord(**base)


def test_columns_match_reference_sheet_exactly():
    """Khuon cot phai KHOP TUYET DOI voi sheet tham chieu - ke ca cot chua bao
    gio dien duoc du lieu. Bo bot cot la loi: ben nhan se phai can chinh tay."""
    assert [header for header, _ in COLUMNS] == REFERENCE_HEADERS


def test_unfilled_reference_columns_are_written_empty_not_dropped(tmp_path):
    path = write_records_to_excel([_record()], tmp_path / "out.xlsx")
    ws = load_workbook(path)["Đèn LED âm trần"]
    header, row = (list(r) for r in ws.iter_rows(max_row=2, values_only=True))

    assert header == REFERENCE_HEADERS
    for column in ("Giá đối chiếu", "VD HDSD", "Mã SAP", "Link mua hàng online"):
        assert row[header.index(column)] is None


def test_one_sheet_per_product_type(tmp_path):
    records = [
        _record(product_id="1", link_san_pham="https://x/1", category_1="Đèn LED âm trần"),
        _record(product_id="2", link_san_pham="https://x/2", category_1="Đèn LED Panel"),
        _record(product_id="3", link_san_pham="https://x/3", category_1="Đèn LED âm trần"),
    ]
    wb = load_workbook(write_records_to_excel(records, tmp_path / "out.xlsx"))

    assert wb.sheetnames == ["Đèn LED âm trần", "Đèn LED Panel"]
    assert wb["Đèn LED âm trần"].max_row == 3  # 1 header + 2 san pham
    assert wb["Đèn LED Panel"].max_row == 2
    # STT dem lai tu 1 trong tung sheet, giong file tham chieu.
    assert [r[0] for r in wb["Đèn LED âm trần"].iter_rows(min_row=2, values_only=True)] == [1, 2]
    assert [r[0] for r in wb["Đèn LED Panel"].iter_rows(min_row=2, values_only=True)] == [1]


def test_records_without_category_go_to_a_review_sheet_at_the_end(tmp_path):
    """Ban ghi fetch loi chua co category van phai duoc giu lai - do chinh la
    ung vien crawl lai o lan chay sau."""
    records = [
        ProductRecord(product_id="1", link_san_pham="https://x/1").mark_error("HTTP 503"),
        _record(product_id="2", link_san_pham="https://x/2"),
    ]
    wb = load_workbook(write_records_to_excel(records, tmp_path / "out.xlsx"))

    assert wb.sheetnames == ["Đèn LED âm trần", UNGROUPED_SHEET]
    assert wb[UNGROUPED_SHEET].max_row == 2


def test_long_sheet_names_keep_the_part_that_tells_them_apart(tmp_path):
    """Excel chan ten sheet > 31 ky tu. Category cua TLC chi khac nhau o DUOI
    ("... Eyecare Pro" / "Standard" / "Premium") nen cat cut duoi la mat thong
    tin: ca 3 ra cung 1 ten roi phai de nhau bang " (2)", " (3)" - nguoi mo file
    khong con biet sheet nao la loai nao.
    """
    variants = ["Pro", "Standard", "Premium"]
    records = [
        _record(
            product_id=str(i), link_san_pham=f"https://x/{i}",
            category_1=f"Đèn LED âm trần thế hệ mới Eyecare {variant}",
        )
        for i, variant in enumerate(variants)
    ]
    wb = load_workbook(write_records_to_excel(records, tmp_path / "out.xlsx"))

    assert len(wb.sheetnames) == 3
    assert all(len(name) <= 31 for name in wb.sheetnames)
    # Phan phan biet phai con nguyen trong ten sheet, va khong sheet nao phai
    # nho toi hau to danh so.
    for variant, name in zip(variants, wb.sheetnames):
        assert name.endswith(variant), name
    assert not any(name.endswith(")") for name in wb.sheetnames)


def test_round_trip_through_all_sheets(tmp_path):
    records = [
        _record(product_id="1", link_san_pham="https://x/1", gia=LIEN_HE),
        _record(product_id="2", link_san_pham="https://x/2", category_1="Đèn LED Panel"),
        ProductRecord(product_id="3", link_san_pham="https://x/3"),
    ]
    path = write_records_to_excel(records, tmp_path / "out.xlsx")

    back = import_legacy_xlsx(path)
    # Doc gop het cac sheet lai - ben goi chi tra cuu theo URL, khong quan tam
    # ban ghi nam sheet nao.
    assert set(back) == {"https://x/1", "https://x/2", "https://x/3"}
    assert back["https://x/1"].gia == LIEN_HE
    assert back["https://x/1"].tags == {"cong_suat": "10W"}
    assert back["https://x/2"].category_1 == "Đèn LED Panel"
    assert back["https://x/3"].crawl_status.value == "partial-missing-fields"


def test_old_file_with_fewer_columns_is_still_readable(tmp_path):
    """File cua phien ban schema cu (17 cot, 1 sheet "Products") phai doc lai
    duoc - neu khong, doi khuon cot dong nghia bat crawl lai tu dau ca site."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    old_headers = [h for h in REFERENCE_HEADERS if h not in ("STT", "Giá đối chiếu", "VD HDSD")]
    ws.append(old_headers)
    row = [None] * len(old_headers)
    row[old_headers.index("Product_ID")] = "1"
    row[old_headers.index("Tên sản phẩm")] = "A"
    row[old_headers.index("Mã Sản Phẩm")] = "X1"
    row[old_headers.index(" category 1 ")] = "Đèn LED âm trần"
    row[old_headers.index("Tags")] = json.dumps({"cong_suat": "10W"}, ensure_ascii=False)
    row[old_headers.index("Giá")] = 166000
    row[old_headers.index("Link sản phẩm")] = "https://x/1"
    row[old_headers.index("Link ảnh sản phẩm")] = "https://x/1.png"
    ws.append(row)
    path = tmp_path / "old.xlsx"
    wb.save(path)

    back = import_legacy_xlsx(path)
    assert back["https://x/1"].crawl_status.value == "ok"
    assert back["https://x/1"].gia_doi_chieu is None
