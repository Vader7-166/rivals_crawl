"""Xuat mot tap chon CAT NGANG nhieu doi thu (capability `excel-export-selection`).

Rang buoc cung cua ca change: khuon 20 cot GIU NGUYEN TUYET DOI. Chieu phan
loai moi (doi thu) chi duoc phep nam o ten file va ten sheet - khong mot cot
nao duoc them, va `category 1` phai giu nguyen van cua site nguon.
"""
import pytest
from openpyxl import load_workbook

from crawler.record import COLUMNS, ProductRecord
from crawler.store import CrawlStore, connect, export_domain, export_selection

KINGLED, TLC = "kingled.com.vn", "tlclighting.com.vn"


class _Fetch:
    def __init__(self, url, html):
        self.url, self.html = url, html
        self.ok, self.status, self.final_url, self.error = True, 200, url, None


@pytest.fixture
def store(tmp_path):
    conn = connect(tmp_path / "crawl.db")
    yield CrawlStore(conn)
    conn.close()


def _add(store, domain, slug, *, category, ten, uu_diem="Tiết kiệm điện"):
    url = f"https://{domain}/{slug}"
    snapshot_id = store.save_snapshot(url, _Fetch(url, f"<html><h1>{ten}</h1></html>"))
    record = ProductRecord(
        product_id="1", link_san_pham=url, ten_san_pham=ten, category_1=category,
        tom_tat_uu_diem_tinh_nang=uu_diem, uu_diem_nguon="keyword",
    )
    store.save_extraction(record, snapshot_id=snapshot_id, extractor_version="v7")
    return url


def _tap_chon(store, **kwargs):
    return {
        KINGLED: [
            _add(store, KINGLED, "dl-1", category="ĐÈN DOWNLIGHT ÂM TRẦN", ten="KL 1"),
            _add(store, KINGLED, "dl-2", category="ĐÈN DOWNLIGHT ÂM TRẦN", ten="KL 2"),
        ],
        TLC: [_add(store, TLC, "am-tran-1/", category="Đèn LED âm trần", ten="TLC 1")],
    }


def test_the_twenty_column_shape_is_untouched(store, tmp_path):
    """6.5: header khop TUYET DOI, ke ca khoang trang thua trong ten cot."""
    chon = _tap_chon(store)

    path, _ = export_selection(store, chon, tmp_path / "ra.xlsx")

    wb = load_workbook(path)
    for ten_sheet in wb.sheetnames:
        assert [c.value for c in wb[ten_sheet][1]] == [h for h, _ in COLUMNS]


def test_one_sheet_per_competitor_for_a_category_scope(store, tmp_path):
    """Pham vi danh muc: nguoi doc dang so hang cua minh voi hang cua TUNG doi
    thu, nen ranh gioi ho can la doi thu."""
    chon = _tap_chon(store)

    path, _ = export_selection(store, chon, tmp_path / "ra.xlsx")

    # Thu tu mac dinh: so san pham giam dan (KingLED 2 > TLC 1).
    assert load_workbook(path).sheetnames == ["KingLED", "TLC Lighting"]


def test_sheet_order_follows_an_explicit_ranking(store, tmp_path):
    """6.3: thu hang la phan doan cua ben nghiep vu nen no chi duoc cham toi
    THU TU SHEET - hang sai thi xuat lai mat 10 giay (design.md muc 3)."""
    chon = _tap_chon(store)

    path, _ = export_selection(store, chon, tmp_path / "ra.xlsx", brand_order=[TLC, KINGLED])

    assert load_workbook(path).sheetnames == ["TLC Lighting", "KingLED"]


def test_a_domain_left_out_of_the_ranking_is_still_exported(store, tmp_path):
    """Bo im lang thi nguoi dung mat du lieu vi mot thu tu hien thi."""
    chon = _tap_chon(store)

    path, _ = export_selection(store, chon, tmp_path / "ra.xlsx", brand_order=[TLC])

    assert set(load_workbook(path).sheetnames) == {"TLC Lighting", "KingLED"}


def test_a_brand_scope_splits_by_product_type_instead(store, tmp_path):
    """Moi ban ghi cung mot doi thu thi chia theo doi thu ra dung MOT sheet -
    khong loi ich gi. Pham vi mot nhan hieu chia theo loai san pham."""
    urls = {KINGLED: [
        _add(store, KINGLED, "a", category="ĐÈN DOWNLIGHT ÂM TRẦN", ten="A"),
        _add(store, KINGLED, "b", category="ĐÈN LED PANEL", ten="B"),
    ]}

    path, _ = export_selection(store, urls, tmp_path / "ra.xlsx", sheet_by_brand=False)

    assert load_workbook(path).sheetnames == ["ĐÈN DOWNLIGHT ÂM TRẦN", "ĐÈN LED PANEL"]


def test_category_1_keeps_the_source_sites_own_words(store, tmp_path):
    """6.6: cot `category 1` mang NGUYEN VAN ten cua site nguon, khong bi thay
    bang tu khoa nguoi dung go. Hai site goi cung mot loai hang hai kieu va ca
    hai kieu deu phai song sot qua file."""
    chon = _tap_chon(store)

    path, _ = export_selection(store, chon, tmp_path / "ra.xlsx")

    wb = load_workbook(path)
    cot = {h: i for i, (h, _) in enumerate(COLUMNS)}[" category 1 "]
    assert [r[cot] for r in wb["KingLED"].iter_rows(min_row=2, values_only=True)] == [
        "ĐÈN DOWNLIGHT ÂM TRẦN", "ĐÈN DOWNLIGHT ÂM TRẦN",
    ]
    assert [r[cot] for r in wb["TLC Lighting"].iter_rows(min_row=2, values_only=True)] == [
        "Đèn LED âm trần",
    ]


def test_a_single_brand_selection_matches_the_per_domain_export(store, tmp_path):
    """6.7: hai duong xuat phai cho ra cung mot thu. Lech nhau tuc la co hai
    khuon file, va khuon thu hai se troi di mot minh."""
    urls = {KINGLED: [
        _add(store, KINGLED, "a", category="ĐÈN DOWNLIGHT ÂM TRẦN", ten="A"),
        _add(store, KINGLED, "b", category="ĐÈN LED PANEL", ten="B"),
    ]}

    theo_domain, _ = export_domain(store, KINGLED, tmp_path / "domain.xlsx")
    theo_chon, _ = export_selection(store, urls, tmp_path / "chon.xlsx", sheet_by_brand=False)

    mot = load_workbook(theo_domain)
    hai = load_workbook(theo_chon)
    assert mot.sheetnames == hai.sheetnames
    for ten in mot.sheetnames:
        assert list(mot[ten].values) == list(hai[ten].values)


def test_the_review_sheet_gathers_rows_across_domains(store, tmp_path):
    """6.4: sheet canh bao hoat dong y nhu duong xuat theo domain, chi la gom
    qua nhieu doi thu."""
    chon = {
        KINGLED: [_add(store, KINGLED, "a", category="C", ten="A", uu_diem=None)],
        TLC: [_add(store, TLC, "b/", category="D", ten="B", uu_diem=None)],
    }

    path, so_dong = export_selection(store, chon, tmp_path / "ra.xlsx")

    assert so_dong == 2
    canh_bao = [s for s in load_workbook(path).sheetnames if s.startswith("⚠")]
    assert len(canh_bao) == 1


def test_a_partially_crawled_selection_still_exports(store, tmp_path):
    """Mot luot crawl la 33-70 phut; bat cho xong moi duoc cam file la bat cho
    vo ich khi phan da co da dung duoc (Open Question 1)."""
    chon = {KINGLED: [
        _add(store, KINGLED, "a", category="C", ten="A"),
        "https://kingled.com.vn/chua-crawl",  # chua co ban ghi nao
    ]}

    path, _ = export_selection(store, chon, tmp_path / "ra.xlsx")

    ten = [r[2] for r in load_workbook(path)["KingLED"].iter_rows(min_row=2, values_only=True)]
    assert ten == ["A"]
