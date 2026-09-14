"""Checkpoint giua chung cua crawl_product_urls (ca duong tuan tu lan song song).

Khong dung mang: fetcher gia tra ve ket qua fetch that bai nen pipeline di
thang toi nhanh tao ban ghi loi - du de kiem tra vong lap va nhip ghi file.
"""
import time
from dataclasses import dataclass
from typing import Optional

import pytest
from openpyxl import load_workbook

from crawler.pipeline import crawl_product_urls


@dataclass
class _FailedFetch:
    ok: bool = False
    html: Optional[str] = None
    status: int = 503
    error: str = "fake"


class _FakeFetcher:
    def __init__(self):
        self.calls = 0

    def fetch(self, url, **kwargs):
        self.calls += 1
        return _FailedFetch()


URLS = [f"https://x/{i}" for i in range(10)]


@pytest.mark.parametrize("workers", [1, 4])
def test_checkpoint_writes_partial_progress(tmp_path, workers):
    path = tmp_path / "out.xlsx"
    records = crawl_product_urls(
        URLS, llm_provider=object(), fetcher=_FakeFetcher(), workers=workers,
        checkpoint_path=path, checkpoint_every=4,
    )

    assert len(records) == len(URLS)
    # Da ghi tam it nhat 1 lan truoc khi chay xong -> file ton tai va doc duoc.
    assert path.exists()
    wb = load_workbook(path)
    written = sum(wb[name].max_row - 1 for name in wb.sheetnames)
    # Lan checkpoint cuoi cung roi vao moc 8/10 san pham.
    assert written == 8


@pytest.mark.parametrize("workers", [1, 4])
def test_no_checkpoint_file_when_path_not_given(tmp_path, workers):
    crawl_product_urls(
        URLS, llm_provider=object(), fetcher=_FakeFetcher(), workers=workers,
    )
    assert not list(tmp_path.iterdir())


def test_reused_records_do_not_cost_a_fetch(tmp_path):
    """Ban ghi da OK tu lan truoc khong duoc fetch lai - day la ca co che
    crawl-lai-co-chon-loc dua tren."""
    from crawler.record import ProductRecord

    prior = ProductRecord(
        product_id="1", ten_san_pham="A", ma_san_pham="X1", category_1="C",
        tags={"a": "b"}, gia=1.0, link_san_pham=URLS[0], link_anh_san_pham="https://i",
    )
    prior.recompute_status()
    fetcher = _FakeFetcher()

    crawl_product_urls(
        URLS, llm_provider=object(), fetcher=fetcher, existing={URLS[0]: prior}, workers=1,
    )
    assert fetcher.calls == len(URLS) - 1


class _SlowFetcher(_FakeFetcher):
    """Ghi lai xem file checkpoint DA ton tai chua tai moi lan fetch."""

    def __init__(self, path):
        super().__init__()
        self._path = path
        self.checkpoint_existed_at_fetch: list[bool] = []

    def fetch(self, url, **kwargs):
        # Nhuong luot cho worker pool chay - khong co no thi phep do "da ghi
        # tam chua" phu thuoc vao thoi diem scheduler nha luong.
        time.sleep(0.05)
        self.checkpoint_existed_at_fetch.append(self._path.exists())
        return super().fetch(url)


def test_checkpoint_is_written_while_fetching_is_still_going(tmp_path):
    """Regression: ban dau phan thu hoach ket qua nam SAU ca vong lap fetch,
    nen checkpoint chi chay o nhung giay cuoi cung cua ca dot - dung luc no vo
    dung. Lan chay TLC 485 san pham khong he co 1 dong "Checkpoint" nao trong
    log vi ly do nay, va dot KingLED 549 san pham (~70 phut) cung se chay suot
    ma khong ghi ra gi neu khong sua.

    Do dung tinh chat can co: den luot fetch cuoi cung thi file checkpoint
    PHAI da ton tai roi.
    """
    path = tmp_path / "out.xlsx"
    fetcher = _SlowFetcher(path)

    crawl_product_urls(
        URLS, llm_provider=object(), fetcher=fetcher, workers=4,
        checkpoint_path=path, checkpoint_every=4,
    )

    assert fetcher.checkpoint_existed_at_fetch[-1], (
        "hết vòng fetch mà vẫn chưa ghi tạm lần nào - checkpoint vô dụng"
    )
