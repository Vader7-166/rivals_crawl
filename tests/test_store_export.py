"""Ket xuat Excel tu kho du lieu: khuon 20 cot giu nguyen, sheet canh bao rieng,
va khong o nao bi cat am tham.
"""
import pytest
from openpyxl import load_workbook

from crawler.record import COLUMNS, REVIEW_HEADERS, REVIEW_SHEET, ProductRecord
from crawler.record.schema import CrawlStatus
from crawler.record.excel_reader import import_legacy_xlsx
from crawler.store import CrawlStore, connect, export_domain, import_xlsx_into_store

PAGE = """
<html><head><title>t</title></head><body>
<nav>menu không thuộc nội dung</nav>
<h1>Đèn LED âm trần 12W</h1>
<div class="desc">
  <h2>4. Ưu điểm nổi bật</h2>
  <ul><li><strong>Tiết kiệm điện:</strong> chỉ 12W.</li>
      <li><strong>Tuổi thọ cao:</strong> 50.000 giờ.</li></ul>
</div>
<script>var rác = 1;</script></body></html>
"""

BARE_PAGE = "<html><body><h1>Đèn không có mục ưu điểm</h1><p>mô tả ngắn.</p></body></html>"


class _Fetch:
    def __init__(self, url, html):
        self.url, self.html = url, html
        self.ok, self.status, self.final_url, self.error = True, 200, url, None


@pytest.fixture
def store(tmp_path):
    conn = connect(tmp_path / "crawl.db")
    yield CrawlStore(conn)
    conn.close()


def _add(store, url, html=PAGE, **kwargs):
    snapshot_id = store.save_snapshot(url, _Fetch(url, html)) if html else None
    kwargs.setdefault("category_1", "Đèn âm trần")
    record = ProductRecord(product_id="1", link_san_pham=url, **kwargs)
    store.save_extraction(record, snapshot_id=snapshot_id, extractor_version="v1")
    return record


def test_the_twenty_column_shape_is_untouched(store, tmp_path):
    _add(store, "https://x.vn/a", ten_san_pham="Đèn A",
         tom_tat_uu_diem_tinh_nang="Tiết kiệm điện", uu_diem_nguon="keyword")
    path, _ = export_domain(store, "x.vn", tmp_path / "out.xlsx")

    ws = load_workbook(path)["Đèn âm trần"]
    assert [c.value for c in ws[1]] == [header for header, _ in COLUMNS]


def test_a_product_without_advantages_keeps_those_cells_empty(store, tmp_path):
    """De trong moi dung - khong duoc nhet text toan trang vao cot uu diem."""
    _add(store, "https://x.vn/a", html=BARE_PAGE, ten_san_pham="Đèn A")
    path, review = export_domain(store, "x.vn", tmp_path / "out.xlsx")

    wb = load_workbook(path)
    ws = wb["Đèn âm trần"]
    headers = [c.value for c in ws[1]]
    row = {h: c.value for h, c in zip(headers, ws[2])}
    assert row["Tóm tắt ưu điểm, tính năng"] is None
    assert row["Nội dung Ưu điểm SP"] is None
    assert review == 1


def test_the_product_also_appears_on_the_review_sheet_with_material_to_work_from(
    store, tmp_path
):
    _add(store, "https://x.vn/den-panel-48w", html=BARE_PAGE, ten_san_pham="Đèn A")
    path, _ = export_domain(store, "x.vn", tmp_path / "out.xlsx")

    ws = load_workbook(path)[REVIEW_SHEET]
    assert [c.value for c in ws[1]] == REVIEW_HEADERS
    url, reason, text, html_file = [c.value for c in ws[2]]
    assert url == "https://x.vn/den-panel-48w"
    assert "Không tìm thấy" in reason
    assert "Đèn không có mục ưu điểm" in text
    assert (tmp_path / html_file).read_text(encoding="utf-8") == BARE_PAGE


def test_page_text_in_the_cell_drops_menu_and_script_noise(store, tmp_path):
    _add(store, "https://x.vn/a", html=BARE_PAGE.replace(
        "<body>", "<body><nav>menu rác</nav><script>var x=1;</script>"))
    path, _ = export_domain(store, "x.vn", tmp_path / "out.xlsx")

    text = load_workbook(path)[REVIEW_SHEET].cell(row=2, column=3).value
    assert "menu rác" not in text and "var x" not in text


def test_a_low_confidence_cell_is_flagged_but_still_delivered(store, tmp_path):
    """Nhanh cum: co du lieu nen sheet danh muc VAN dien, nhung len sheet canh
    bao de soi lai - day dung la nhom 6/360 o lay nham cua KingLED."""
    _add(store, "https://x.vn/a", ten_san_pham="Đèn A",
         tom_tat_uu_diem_tinh_nang="Thanh nhôm\nBộ nguồn", uu_diem_nguon="cluster")
    path, review = export_domain(store, "x.vn", tmp_path / "out.xlsx")

    wb = load_workbook(path)
    assert wb["Đèn âm trần"].cell(row=2, column=16).value == "Thanh nhôm\nBộ nguồn"
    assert review == 1
    assert "nhánh cụm" in wb[REVIEW_SHEET].cell(row=2, column=2).value


def test_a_row_imported_from_xlsx_has_no_attachment_but_still_shows_up(store, tmp_path):
    _add(store, "https://x.vn/cu", html=None, ten_san_pham="Đèn cũ")
    path, review = export_domain(store, "x.vn", tmp_path / "out.xlsx")

    url, reason, text, html_file = [
        c.value for c in load_workbook(path)[REVIEW_SHEET][2]
    ]
    assert review == 1
    assert "Chưa có HTML" in reason
    assert text is None and html_file is None


def test_no_cell_is_ever_silently_truncated(store, tmp_path):
    """openpyxl ghi vuot 32.767 ky tu ma khong bao gi, doc lai chi con 32.767 -
    o cat cut trong y het o lanh. Moi o phai di qua bo chan CO BAO."""
    _add(store, "https://x.vn/a", ten_san_pham="Đèn A",
         noi_dung_uu_diem_sp="x" * 40_000, uu_diem_nguon="keyword")
    path, _ = export_domain(store, "x.vn", tmp_path / "out.xlsx")

    wb = load_workbook(path)
    for sheet in wb.sheetnames:
        for row in wb[sheet].iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    assert len(value) <= 32_767
    cell = wb["Đèn âm trần"].cell(row=2, column=20).value
    assert "đã cắt" in cell  # dau hieu nhin thay duoc, khong cat im lang


def test_a_page_whose_text_is_enormous_does_not_break_the_export(store, tmp_path):
    """Ca ngoai le da do: `kingled_product.html` (ban chua render) cho ra
    691.236 ky tu text vi mang mot khoi du lieu lon khong nam trong <script>."""
    huge = "<html><body><h1>Đèn</h1><p>" + ("dữ liệu " * 100_000) + "</p></body></html>"
    _add(store, "https://x.vn/a", html=huge)
    path, _ = export_domain(store, "x.vn", tmp_path / "out.xlsx")

    text = load_workbook(path)[REVIEW_SHEET].cell(row=2, column=3).value
    assert len(text) <= 32_767
    assert "đã cắt" in text


def test_import_skips_the_review_sheet_without_being_told_to(store, tmp_path):
    """Doc lai file da xuat: dong cua sheet canh bao KHONG duoc thanh san pham.

    `import_legacy_xlsx` gop PHANG moi sheet theo URL nen khong co ranh gioi
    nao khac de phan biet - vi vay sheet nay bi loai MAC DINH, khong doi ben
    goi tu khai bao.
    """
    _add(store, "https://x.vn/a", html=BARE_PAGE, ten_san_pham="Đèn A")
    path, _ = export_domain(store, "x.vn", tmp_path / "out.xlsx")
    assert REVIEW_SHEET in load_workbook(path).sheetnames

    back = import_legacy_xlsx(path)
    assert list(back) == ["https://x.vn/a"]
    assert back["https://x.vn/a"].ten_san_pham == "Đèn A"


def test_round_trip_through_the_store_preserves_every_cell(store, tmp_path):
    """Nhap 1 file .xlsx vao kho roi xuat lai -> khop tung o tren sheet san pham.

    Day la phep so bien "khuon dau ra khong doi" tu mot loi hua thanh mot kiem
    chung (task 9.1 cua change nay lam dieu tuong tu tren du lieu that).
    """
    from crawler.record import write_records_to_excel

    goc = tmp_path / "goc.xlsx"
    write_records_to_excel(
        [
            ProductRecord(
                product_id="p1", link_san_pham="https://x.vn/a", ten_san_pham="Đèn A",
                ma_san_pham="A1", category_1="Đèn âm trần", gia=339900.0,
                link_anh_san_pham="https://x.vn/a.jpg",
                tags={"cong_suat": "12W", "anh_sang": "3 màu"},
                tom_tat_uu_diem_tinh_nang="Tiết kiệm điện\nTuổi thọ cao",
                noi_dung_uu_diem_sp="Ưu điểm nổi bật\nTiết kiệm điện",
                thong_so_ky_thuat="Công suất: 12W",
            ),
            ProductRecord(
                product_id="p2", link_san_pham="https://x.vn/b/", ten_san_pham="Đèn B",
                category_1="Đèn panel", gia="Liên hệ",
            ),
        ],
        goc,
    )

    import_xlsx_into_store(goc, store)
    lai, _ = export_domain(store, "x.vn", tmp_path / "lai.xlsx")

    truoc, sau = load_workbook(goc), load_workbook(lai)
    # Sheet canh bao la phan THEM VAO co chu dich; phep so la tren cac sheet
    # SAN PHAM - do moi la thu phai khop tuyet doi voi khuon tham chieu.
    assert set(truoc.sheetnames) == set(sau.sheetnames) - {REVIEW_SHEET}
    for name in truoc.sheetnames:
        a = [row for row in truoc[name].iter_rows(values_only=True)]
        b = [row for row in sau[name].iter_rows(values_only=True)]
        assert a == b, f"sheet '{name}' lệch"


# -- Trang khong phai san pham ----------------------------------------------


def test_a_non_product_page_stays_out_of_the_file(store, tmp_path):
    """68 ban ghi cua Roman/Duhal la BAI VIET chu khong phai san pham. Chung
    khong duoc di vao file ket qua - nhung van nam nguyen trong kho, nen doi y
    thi chi phai xuat lai, khong phai crawl lai."""
    _add(store, "https://x.vn/den-a", ten_san_pham="Đèn A")
    url = "https://x.vn/10-thong-so-den-led"
    bai_viet = _add(store, url, ten_san_pham="10 thông số")
    bai_viet.mark_not_a_product("Thông tin")
    store.save_extraction(
        bai_viet, snapshot_id=store.snapshot_id_for(url), extractor_version="v7"
    )

    path, _ = export_domain(store, "x.vn", tmp_path / "out.xlsx")

    names = [row[2] for row in load_workbook(path)["Đèn âm trần"].iter_rows(
        min_row=2, values_only=True)]
    assert names == ["Đèn A"]
    assert "https://x.vn/10-thong-so-den-led" in store.current_records("x.vn")


def test_a_non_product_is_never_queued_for_crawling_again(store):
    """Crawl lai bao nhieu lan cung ra dung ket luan do. De no trong hang doi
    la fetch lai 68 URL o MOI luot chay."""
    url = "https://x.vn/bai-viet"
    record = _add(store, url, ten_san_pham="Bài viết")
    record.mark_not_a_product("Tin tức")
    store.save_extraction(
        record, snapshot_id=store.snapshot_id_for(url), extractor_version="v7"
    )

    assert store.urls_needing_crawl([url]) == []


def test_a_page_missing_fields_is_still_queued(store):
    """Doi chung: NOT_A_PRODUCT khac han "san pham that nhung thieu o" - cai
    sau VAN phai duoc crawl lai."""
    url = "https://x.vn/thieu-o"
    record = _add(store, url, ten_san_pham=None)
    record.recompute_status()
    store.save_extraction(
        record, snapshot_id=store.snapshot_id_for(url), extractor_version="v7"
    )

    assert record.crawl_status is CrawlStatus.PARTIAL_MISSING_FIELDS
    assert store.urls_needing_crawl([url]) == [url]
