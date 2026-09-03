"""Doc lai file .xlsx da xuat truoc day de NHAP vao kho du lieu.

Truoc day day la nguon trang thai cua co che crawl-lai-co-chon-loc; vai tro do
nay thuoc ve kho du lieu (`store.CrawlStore.current_records`), von khong phu
thuoc vao su ton tai cua file .xlsx nao. Ham o day doi ten tu
`load_existing_records` thanh `import_legacy_xlsx` chinh de khong con doc nhu
mot API trang thai dung chung - no la duong DI MOT CHIEU tu dinh dang cu vao
kho.

Vi field trang thai crawl khong duoc luu trong file Excel (xem excel_writer),
trang thai duoc SUY RA LAI khi doc: ban ghi thieu 1 trong cac field bat buoc
(REQUIRED_FIELDS) duoc coi la ung vien crawl lai, bat ke ly do ban dau la loi
fetch hay thieu field khi trich xuat.

Cot duoc ghep theo TEN header (da strip) chu khong theo vi tri - xem chu thich
trong import_legacy_xlsx.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .excel_writer import REVIEW_SHEET
from .price import normalize_price
from .schema import COLUMNS, ProductRecord


def import_legacy_xlsx(
    input_path: str | Path,
    skip_sheets: Iterable[str] | None = None,
) -> dict[str, ProductRecord]:
    """Doc file .xlsx da co, tra ve map link_san_pham -> ProductRecord (da
    recompute_status()). Tra ve dict rong neu file chua ton tai.

    File duoc ghi theo kieu MOI LOAI SAN PHAM 1 SHEET (xem excel_writer) nen o
    day gop toan bo sheet lai thanh 1 map phang - ben goi chi can tra cuu theo
    URL de biet ban ghi nao da OK, khong quan tam no nam sheet nao.
    """
    input_path = Path(input_path)
    if not input_path.exists():
        return {}

    wb = load_workbook(input_path, read_only=True, data_only=True)

    # Sheet canh bao bi bo qua MAC DINH: no do chinh buoc ket xuat sinh ra,
    # khong chua san pham, va khong co cot `Product_ID` nen doc vao se lam
    # ProductRecord no thang. Bat moi ben goi tu nho la mot cai bay khong can
    # thiet. `skip_sheets` de them ten khac.
    skip = {REVIEW_SHEET} | set(skip_sheets or ())
    records: dict[str, ProductRecord] = {}
    for sheet_name in wb.sheetnames:
        # Sheet canh bao (do buoc ket xuat sinh ra) KHONG chua san pham - doc no
        # vao day se bien cac dong "viec can lam" thanh ban ghi san pham. Ham
        # nay gop PHANG moi sheet theo URL nen khong co ranh gioi nao khac de
        # phan biet, phai loai bang ten.
        if sheet_name in skip:
            continue
        _read_sheet_into(wb[sheet_name], records)
    return records


def _read_sheet_into(ws, records: dict[str, ProductRecord]) -> None:
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return

    # Doi chieu theo TEN cot chu khong theo vi tri, va bo qua khoang trang thua
    # trong header (khuon tham chieu co san ' category 1 ' / 'Thong so ky thuat ').
    # Nho vay 1 file cua phien ban schema cu - thieu vai cot moi them - van doc
    # lai duoc de crawl-lai-co-chon-loc, thay vi bi tu choi toan bo va bat crawl
    # lai tu dau; cot thieu chi don gian thanh None.
    field_by_header = {h.strip(): f for h, f in COLUMNS}
    column_fields: list[str | None] = [
        field_by_header.get(str(h).strip()) if h is not None else None for h in header
    ]
    if "link_san_pham" not in column_fields:
        raise ValueError(
            f"Sheet '{ws.title}' khong co cot 'Link sản phẩm' - khong xac dinh "
            "duoc ban ghi ung voi URL nao, khong dung lam du lieu cu duoc."
        )

    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        values = {
            field_name: value
            for field_name, value in zip(column_fields, row)
            if field_name is not None
        }
        # STT bi bo di co chu dich: no la so thu tu DONG trong sheet cu. Neu
        # giu lai, ban ghi doi sheet (category doi) hoac sheet bi chen them dong
        # se lam cot STT thung/trung so o lan ghi sau. Cu de writer danh lai.
        values["stt"] = None
        values["tags"] = json.loads(values["tags"]) if values.get("tags") else {}
        values["gia"] = normalize_price(values.get("gia"))
        values["gia_doi_chieu"] = normalize_price(values.get("gia_doi_chieu"))
        record = ProductRecord(**values)
        record.recompute_status()
        key = record.link_san_pham or record.product_id
        records[key] = record


def records_needing_recrawl(records: dict[str, ProductRecord]) -> list[ProductRecord]:
    """Cac ban ghi co trang thai khac OK - ung vien cho lan crawl lai tiep theo."""
    from .schema import CrawlStatus

    return [r for r in records.values() if r.crawl_status != CrawlStatus.OK]


__all__ = ["import_legacy_xlsx", "records_needing_recrawl"]
