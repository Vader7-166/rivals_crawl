"""Ghi danh sach ProductRecord ra 1 file .xlsx rieng cho 1 site. Task 2.3.

Cau truc file bam theo chinh product_Metadata (1).xlsx: MOT file cho ca site,
chia thanh NHIEU SHEET, moi sheet la 1 loai san pham. Loai san pham lay tu
`category 1` cua chinh site nguon (khong anh xa sang taxonomy khac - xem spec
product-record-schema, muc "Category giu nguyen theo site nguon").

Moi sheet ghi DU 20 cot cua sheet "2. LED Downlight" trong file tham chieu,
dung ten va dung thu tu. Cot nao site doi thu khong co du lieu tuong duong thi
o do de RONG - khong duoc bo cot ra khoi file, vi khuon cot lech thi ben nhan
phai can chinh tay truoc khi ghep vao file tham chieu.

Field trang thai crawl noi bo (crawl_status/crawl_error) van KHONG xuat ra cot
Excel nao - do la du lieu van hanh, khong thuoc khuon tham chieu (xem
specs/product-record-schema/spec.md).
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Iterable, Optional, Sequence

from openpyxl import Workbook

from .price import LIEN_HE
from .schema import COLUMNS, ProductRecord

# Sheet cho cac ban ghi khong xac dinh duoc loai (category 1 rong - thuong la
# ban ghi fetch loi). Gom rieng thay vi vut di de con crawl lai duoc.
UNGROUPED_SHEET = "Chưa phân loại"

# Sheet liet ke san pham can nguoi xu ly tay. Nam NGOAI khuon 20 cot: them
# COT moi thi pha vo hop dong "khop tuyet doi voi khuon tham chieu", con nhet
# text toan trang vao cot `Nội dung Ưu điểm SP` thi ben nhan khong phan biet
# duoc "day la muc uu diem that" voi "day la ca trang, tu tim lay" - dung loai
# loi da tung gap (bang thong so bi ghi vao cot uu diem). Nen: mot sheet rieng.
REVIEW_SHEET = "⚠ Cần xử lý tay"
REVIEW_HEADERS = ["Link sản phẩm", "Lý do", "Toàn văn text trang", "File HTML"]

# Gioi han cung cua Excel cho 1 o. openpyxl KHONG kiem tra: no ghi "thanh cong"
# khong mot canh bao nao, roi doc lai chi con 32.767 ky tu - o cat cut trong y
# het o lanh. Do that: HTML mot trang KingLED la 146.199 ky tu, mat 78% trong
# im lang. Nen moi o deu di qua `_fit_cell` truoc khi ghi.
_EXCEL_CELL_LIMIT = 32_767
_CELL_TRUNCATE_AT = 32_000
_TRUNCATED_MARK = "\n[… đã cắt ở 32.000 ký tự — bản đầy đủ ở kho dữ liệu / file .html đính kèm]"

# Excel: ten sheet toi da 31 ky tu va cam : \ / ? * [ ]
_INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")
_MAX_SHEET_NAME = 31


def _fit_cell(value):
    """Cat CO BAO thay vi de openpyxl cat am tham. Ap cho MOI o cua moi sheet."""
    if not isinstance(value, str) or len(value) <= _EXCEL_CELL_LIMIT:
        return value
    return value[:_CELL_TRUNCATE_AT] + _TRUNCATED_MARK


def _cell_value_for(record: ProductRecord, field_name: str, row_number: int):
    if field_name == "stt":
        # So thu tu la thuoc tinh cua DONG chu khong phai cua san pham: neu ban
        # ghi chua co thi danh theo vi tri dong de cot khong bao gio thung.
        return record.stt if record.stt is not None else row_number
    if field_name == "tags":
        return json.dumps(record.tags, ensure_ascii=False) if record.tags else None
    if field_name == "gia":
        value = record.gia
        return value  # float, LIEN_HE (str), hoac None deu ghi thang duoc
    return getattr(record, field_name)


def _shorten(name: str) -> str:
    """Cat ten cho vua gioi han cua Excel, giu CA DAU LAN DUOI.

    Cat cut duoi (`name[:31]`) la cach lam mat thong tin nhat voi chinh du lieu
    TLC: "Đèn LED âm trần thế hệ mới Eyecare Pro / Standard / Premium" chi khac
    nhau o DUOI, cat cut duoi thi ca 3 loai ra cung 1 ten
    "Đèn LED âm trần thế hệ mới Eyec" roi phai de nhau bang hau to " (2)",
    " (3)" - nguoi mo file khong con biet sheet nao la loai nao.
    """
    if len(name) <= _MAX_SHEET_NAME:
        return name
    head = _MAX_SHEET_NAME // 2
    tail = _MAX_SHEET_NAME - head - 1  # 1 ky tu danh cho dau "…"
    return name[:head] + "…" + name[-tail:]


def _sheet_name(raw: str, taken: set[str]) -> str:
    """Ten sheet hop le, duy nhat, giu duoc cang nhieu ten goc cang tot."""
    name = _shorten(_INVALID_SHEET_CHARS.sub("-", raw).strip() or UNGROUPED_SHEET)
    if name not in taken:
        return name
    # Van trung sau khi cat giua (2 category giong nhau ca dau lan duoi): danh
    # so de khong nhom nao bi de len nhom khac va mat sach du lieu.
    for suffix_index in range(2, 100):
        suffix = f" ({suffix_index})"
        candidate = name[: _MAX_SHEET_NAME - len(suffix)] + suffix
        if candidate not in taken:
            return candidate
    raise ValueError(f"Khong sinh duoc ten sheet duy nhat cho '{raw}'")


def group_records_by_type(records: Iterable[ProductRecord]) -> dict[str, list[ProductRecord]]:
    """Gom ban ghi theo loai san pham (`category 1` cua site nguon).

    Giu THU TU XUAT HIEN dau tien cua moi loai thay vi sap xep A-Z: thu tu do
    phan anh thu tu duyet danh muc cua site, on dinh giua cac lan chay va de
    doi chieu voi site hon la thu tu bang chu cai.
    """
    grouped: dict[str, list[ProductRecord]] = {}
    for record in records:
        key = (record.category_1 or "").strip() or UNGROUPED_SHEET
        grouped.setdefault(key, []).append(record)
    # Nhom "chua phan loai" luon xuong cuoi file - do la phan can review, khong
    # nen chen giua cac sheet du lieu that.
    if UNGROUPED_SHEET in grouped and len(grouped) > 1:
        grouped[UNGROUPED_SHEET] = grouped.pop(UNGROUPED_SHEET)
    return grouped


def write_records_to_excel(
    records: Iterable[ProductRecord],
    output_path: str | Path,
    review_rows: Optional[Iterable[Sequence]] = None,
) -> Path:
    """Ghi toan bo ban ghi cua 1 site ra 1 file .xlsx, moi loai san pham 1 sheet.

    `review_rows` (tuy chon) la cac dong cua sheet canh bao `REVIEW_SHEET`,
    moi dong theo dung thu tu `REVIEW_HEADERS`. San pham can xu ly tay xuat
    hien o CA HAI cho: tren sheet danh muc (2 o uu diem de TRONG - de trong moi
    dung, khong bia du lieu) va tren sheet canh bao kem nguyen lieu de xu ly.
    Khong truyen gi thi file khong co sheet canh bao, y nhu truoc.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    grouped = group_records_by_type(records)
    headers = [header for header, _ in COLUMNS]

    wb = Workbook()
    wb.remove(wb.active)  # bo sheet mac dinh, chi giu sheet do minh tao

    taken: set[str] = set()
    for group_name, group_records in grouped.items():
        name = _sheet_name(group_name, taken)
        taken.add(name)

        ws = wb.create_sheet(title=name)
        ws.append(headers)
        # STT dem lai tu 1 trong TUNG sheet - giong file tham chieu, moi sheet
        # la 1 bang doc lap.
        for row_number, record in enumerate(group_records, start=1):
            ws.append(
                [
                    _fit_cell(_cell_value_for(record, field_name, row_number))
                    for _, field_name in COLUMNS
                ]
            )

        for col_idx, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(
                12, min(60, len(header) + 4)
            )

    if not wb.sheetnames:  # khong co ban ghi nao - van tao file co khuon cot
        ws = wb.create_sheet(title=UNGROUPED_SHEET)
        ws.append(headers)

    rows = list(review_rows or ())
    if rows:
        ws = wb.create_sheet(title=REVIEW_SHEET)
        ws.append(REVIEW_HEADERS)
        for row in rows:
            ws.append([_fit_cell(value) for value in row])
        for col_idx, width in enumerate((55, 40, 90, 40), start=1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    wb.save(output_path)
    return output_path


__all__ = [
    "write_records_to_excel",
    "group_records_by_type",
    "UNGROUPED_SHEET",
    "REVIEW_SHEET",
    "REVIEW_HEADERS",
    "LIEN_HE",
]
