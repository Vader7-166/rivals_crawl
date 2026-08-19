"""Do dac thong so cua 1 file .xlsx da crawl xong - CHAT LUONG, khong phai toc do.

Toc do do o scripts/tune_crawl.py va o chinh log lan chay; script nay tra loi
cau hoi khac: du lieu lay ve co dung va co day khong.

Chay: .venv/bin/python scripts/report_crawl.py [duong_dan_file.xlsx]
"""
from __future__ import annotations

import json
import re
import statistics
import sys
import unicodedata
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from openpyxl import load_workbook  # noqa: E402

from crawler.record import COLUMNS, load_existing_records  # noqa: E402
from crawler.record.price import LIEN_HE  # noqa: E402

DEFAULT_PATH = Path("output/tlclighting_all.xlsx")


def _norm(s: str) -> str:
    return re.sub(r"[\s ]+", "", unicodedata.normalize("NFC", str(s)).lower())


def grounding_rate(records) -> tuple[float, int]:
    """Ty le gia tri tag thuc su xuat hien trong text nguon cua chinh trang do.

    Thuoc do CHONG BIA: LLM co the tra ve tag dep de nhung bia ra thi gia tri
    se khong co trong nguon. Doi chieu voi `Thong so ky thuat` + ten san pham -
    nguon nay khong hoan hao (co san pham bo trong cot do) nen doc nhu CAN
    DUOI cua do trung thuc, khong phai ty le tuyet doi.
    """
    total = hit = 0
    for r in records:
        source = _norm((r.thong_so_ky_thuat or "") + " " + (r.ten_san_pham or ""))
        if not source:
            continue
        for value in (r.tags or {}).values():
            total += 1
            pieces = [p for p in re.split(r"[/,;()]+", _norm(value)) if p]
            if pieces and all(p in source for p in pieces):
                hit += 1
    return (hit / total if total else 0.0), total


def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_PATH
    if not path.exists():
        raise SystemExit(f"Không thấy file {path}")

    wb = load_workbook(path, read_only=True)
    sheets = {name: wb[name].max_row - 1 for name in wb.sheetnames}
    records = list(load_existing_records(path).values())
    total = len(records)

    print(f"FILE: {path}   ({total} bản ghi, {len(sheets)} sheet)\n")

    print("PHÂN BỐ THEO LOẠI SẢN PHẨM (mỗi sheet 1 loại)")
    for name, count in sorted(sheets.items(), key=lambda kv: -kv[1]):
        print(f"  {count:>4}  {name}")

    print("\nĐỘ ĐẦY CỦA TỪNG CỘT (số ô có dữ liệu / tổng)")
    for header, field_name in COLUMNS:
        if field_name == "stt":
            continue
        filled = sum(
            1 for r in records
            if getattr(r, field_name) not in (None, "", {})
        )
        bar = "█" * round(filled / total * 20) if total else ""
        print(f"  {header.strip():<28} {filled:>4}/{total:<4} {filled/total*100 if total else 0:>5.1f}%  {bar}")

    numeric = sum(1 for r in records if isinstance(r.gia, float))
    contact = sum(1 for r in records if r.gia == LIEN_HE)
    print("\nGIÁ (3 trạng thái, không được gộp lẫn)")
    print(f"  số cụ thể   {numeric:>4}  ({numeric/total*100 if total else 0:.1f}%)")
    print(f"  \"Liên hệ\"   {contact:>4}  ({contact/total*100 if total else 0:.1f}%)")
    print(f"  trống       {total-numeric-contact:>4}  (chưa crawl được → ứng viên crawl lại)")
    prices = [r.gia for r in records if isinstance(r.gia, float)]
    if prices:
        print(f"  khoảng giá  {min(prices):,.0f} – {max(prices):,.0f} đ "
              f"(trung vị {statistics.median(prices):,.0f} đ)")

    tag_counts = [len(r.tags or {}) for r in records]
    rate, tag_total = grounding_rate(records)
    keys = Counter(k for r in records for k in (r.tags or {}))
    print("\nTAGS (tầng 2 - LLM)")
    print(f"  sản phẩm có tag       {sum(1 for c in tag_counts if c):>4}/{total}")
    print(f"  số tag trung bình     {statistics.mean(tag_counts) if tag_counts else 0:>6.1f} "
          f"(ít nhất {min(tag_counts) if tag_counts else 0}, nhiều nhất {max(tag_counts) if tag_counts else 0})")
    print(f"  tỷ lệ tag CÓ CĂN CỨ   {rate*100:>6.1f}%  ({tag_total} giá trị được đối chiếu)")
    print("  thuộc tính hay gặp:   " + ", ".join(f"{k} ({c})" for k, c in keys.most_common(8)))

    needs_review = [r for r in records if r.crawl_status.value != "ok"]
    print(f"\nTRẠNG THÁI: {total - len(needs_review)}/{total} OK, {len(needs_review)} cần review")
    if needs_review:
        reasons = Counter(
            ", ".join(r.missing_required_fields()) or "lỗi fetch" for r in needs_review
        )
        for reason, count in reasons.most_common(10):
            print(f"  {count:>4}  thiếu: {reason}")
        print("  URL cần crawl lại (tối đa 10):")
        for r in needs_review[:10]:
            print(f"    {r.link_san_pham}")


if __name__ == "__main__":
    main()
